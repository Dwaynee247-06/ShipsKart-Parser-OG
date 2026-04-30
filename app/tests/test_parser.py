"""Unit tests for the Excel parser service."""
from pathlib import Path
import openpyxl
import pytest

from app.services.parsers.excel import extract_workbook, _normalize_header, _find_blocks


def make_sample_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Table 1 — product list (matches example from screenshot)
    headers = ["Sr No", "Product name", "Unit of measurement", "quantity", "Remarks"]
    ws.append(headers)
    ws.append([1, "A", "KG", 1, None])
    ws.append([2, "B", "LTR", 2, None])
    ws.append([3, "C", "GPH", 4, None])
    ws.append([4, "D", "POUND", 5, None])
    ws.append([5, "E", "HA", 6, None])
    ws.append([6, "F", "CADSB", 35, None])
    wb.save(path)


def test_normalize_header_aliases():
    assert _normalize_header("Sr No") == "sr_no"
    assert _normalize_header("Unit of measurement") == "unit_of_measurement"
    assert _normalize_header("Qty") == "quantity"
    assert _normalize_header("Remarks") == "remarks"
    assert _normalize_header(None) == ""


def test_extract_workbook(tmp_path):
    xlsx = tmp_path / "sample.xlsx"
    make_sample_xlsx(xlsx)
    result = extract_workbook(xlsx)

    assert "requirements1" in result
    req = result["requirements1"]
    assert req.total_rows == 6
    assert "sr_no" in req.headers
    assert req.tables[0].sheet_name == "Sheet1"

    # Blank cells preserved as None
    first_row = req.tables[0].rows[0]
    assert first_row["remarks"] is None
